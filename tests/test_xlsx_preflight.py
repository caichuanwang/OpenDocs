from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

import opendocs.parsers.xlsx.preflight as preflight_module
from opendocs.errors import CorruptDocumentError, LimitExceededError, UnsupportedDocumentError
from opendocs.options import ParseOptions
from opendocs.parsers.xlsx import XlsxParser
from opendocs.parsers.xlsx.models import XlsxSheetKind, XlsxSheetState
from opendocs.parsers.xlsx.preflight import MAX_SHEETS, preflight_xlsx
from opendocs.source import ResolvedSource
from tests.xlsx_fixtures import rewrite_xlsx, write_structured_xlsx

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _one_sheet(path: Path, *, cells: tuple[str, ...] = ()) -> None:
    write_structured_xlsx(
        path,
        sheets=(("Sheet", "worksheet", "visible", "A1:B2", cells),),
    )


def _worksheet(body: str) -> bytes:
    return (f'<worksheet xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}">{body}</worksheet>').encode()


def _relationships(*items: tuple[str, str, str]) -> bytes:
    values = "".join(
        f'<Relationship Id="{relationship_id}" Type="{relationship_type}" Target="{target}"/>'
        for relationship_id, relationship_type, target in items
    )
    return f'<Relationships xmlns="{PACKAGE_REL_NS}">{values}</Relationships>'.encode()


def test_preflight_preserves_worksheet_chartsheet_state_and_empty_sheet_order(
    tmp_path: Path,
) -> None:
    path = tmp_path / "ordered.xlsx"
    write_structured_xlsx(
        path,
        sheets=(
            ("Visible", "worksheet", "visible", "A1", ("A1",)),
            ("Hidden", "worksheet", "hidden", None, ()),
            ("Very Hidden", "worksheet", "veryHidden", "C3", ("C3",)),
            ("Chart", "chartsheet", "visible", None, ()),
        ),
    )

    result = preflight_xlsx(path)

    assert [(sheet.sheet_index, sheet.name) for sheet in result.sheets] == [
        (1, "Visible"),
        (2, "Hidden"),
        (3, "Very Hidden"),
        (4, "Chart"),
    ]
    assert [sheet.kind for sheet in result.sheets] == [
        XlsxSheetKind.WORKSHEET,
        XlsxSheetKind.WORKSHEET,
        XlsxSheetKind.WORKSHEET,
        XlsxSheetKind.CHARTSHEET,
    ]
    assert [sheet.state for sheet in result.sheets] == [
        XlsxSheetState.VISIBLE,
        XlsxSheetState.HIDDEN,
        XlsxSheetState.VERY_HIDDEN,
        XlsxSheetState.VISIBLE,
    ]
    assert result.serialized_cells == 2


def test_preflight_accepts_openpyxl_package_absolute_relationship_targets(
    tmp_path: Path,
) -> None:
    path = tmp_path / "openpyxl.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    chart = BarChart()
    chart.add_data(Reference(workbook.active, min_col=1, min_row=1, max_row=1))
    chart_sheet = workbook.create_chartsheet("Chart")
    chart_sheet.add_chart(chart)
    workbook.save(path)

    result = preflight_xlsx(path)

    assert [(sheet.name, sheet.kind) for sheet in result.sheets] == [
        ("Sheet", XlsxSheetKind.WORKSHEET),
        ("Chart", XlsxSheetKind.CHARTSHEET),
    ]


def test_preflight_accepts_128_sheets_and_rejects_129(tmp_path: Path) -> None:
    accepted = tmp_path / "accepted.xlsx"
    rejected = tmp_path / "rejected.xlsx"
    sheets = tuple(
        (f"Sheet {index}", "worksheet", "visible", None, ()) for index in range(1, MAX_SHEETS + 1)
    )
    write_structured_xlsx(accepted, sheets=sheets)
    write_structured_xlsx(
        rejected,
        sheets=(*sheets, ("One too many", "worksheet", "visible", None, ())),
    )

    assert len(preflight_xlsx(accepted).sheets) == MAX_SHEETS
    with pytest.raises(LimitExceededError, match="sheet count"):
        preflight_xlsx(rejected)


def test_sparse_full_grid_dimension_fails_before_any_loader(tmp_path: Path) -> None:
    path = tmp_path / "sparse.xlsx"
    write_structured_xlsx(
        path,
        sheets=(("Sparse", "worksheet", "visible", "A1:XFD1048576", ("A1",)),),
    )

    with pytest.raises(LimitExceededError, match="declared dimension"):
        preflight_xlsx(path)


@pytest.mark.asyncio
async def test_parser_seam_preflights_limits_and_does_not_map_max_pages_to_sheets(
    tmp_path: Path,
) -> None:
    parser = XlsxParser()
    accepted = tmp_path / "two-sheets.xlsx"
    rejected = tmp_path / "too-many-sheets.xlsx"
    write_structured_xlsx(
        accepted,
        sheets=(
            ("One", "worksheet", "visible", None, ()),
            ("Two", "worksheet", "visible", None, ()),
        ),
    )
    write_structured_xlsx(
        rejected,
        sheets=tuple(
            (f"S{index}", "worksheet", "visible", None, ()) for index in range(MAX_SHEETS + 1)
        ),
    )

    with pytest.raises(UnsupportedDocumentError, match="content parsing"):
        await parser.parse(
            ResolvedSource(accepted, "two-sheets.xlsx", False),
            options=ParseOptions(max_pages=1),
        )
    with pytest.raises(LimitExceededError, match="sheet count"):
        await parser.parse(
            ResolvedSource(rejected, "too-many-sheets.xlsx", False),
            options=ParseOptions(max_pages=1),
        )


@pytest.mark.parametrize(
    ("limit_name", "body", "message"),
    [
        (
            "MAX_SERIALIZED_CELLS",
            '<dimension ref="A1:B1"/><sheetData><row><c r="A1"/><c r="B1"/></row></sheetData>',
            "serialized cell",
        ),
        (
            "MAX_NON_EMPTY_CELLS",
            '<dimension ref="A1:B1"/><sheetData><row>'
            '<c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row></sheetData>',
            "non-empty cell",
        ),
        (
            "MAX_MERGE_RANGES",
            '<dimension ref="A1:D1"/><sheetData/>'
            '<mergeCells><mergeCell ref="A1:B1"/><mergeCell ref="C1:D1"/></mergeCells>',
            "merge range",
        ),
        (
            "MAX_MERGE_FOOTPRINT",
            '<dimension ref="A1:C1"/><sheetData/><mergeCells><mergeCell ref="A1:C1"/></mergeCells>',
            "merge footprint",
        ),
        (
            "MAX_CONDITIONAL_FORMATTING_RULES",
            '<dimension ref="A1"/><sheetData/>'
            '<conditionalFormatting sqref="A1"><cfRule type="expression" priority="1"/>'
            '<cfRule type="expression" priority="2"/></conditionalFormatting>',
            "conditional formatting",
        ),
        (
            "MAX_DATA_VALIDATIONS",
            '<dimension ref="A1"/><sheetData/><dataValidations count="2">'
            '<dataValidation type="whole" sqref="A1"/>'
            '<dataValidation type="whole" sqref="A1"/></dataValidations>',
            "data validation",
        ),
        (
            "MAX_ROW_DIMENSIONS",
            '<dimension ref="A1:A2"/><sheetData><row r="1" hidden="1"/>'
            '<row r="2" hidden="1"/></sheetData>',
            "row dimension",
        ),
        (
            "MAX_COLUMN_DIMENSIONS",
            '<dimension ref="A1:B1"/><cols><col min="1" max="1" width="10"/>'
            '<col min="2" max="2" width="10"/></cols><sheetData/>',
            "column dimension",
        ),
        (
            "MAX_PAGE_BREAKS",
            '<dimension ref="A1"/><sheetData/><rowBreaks count="2">'
            '<brk id="1"/><brk id="2"/></rowBreaks>',
            "page break",
        ),
        (
            "MAX_SCENARIOS",
            '<dimension ref="A1"/><sheetData/><scenarios>'
            '<scenario name="one"/><scenario name="two"/></scenarios>',
            "scenario",
        ),
    ],
)
def test_worksheet_loader_collections_are_bounded(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    limit_name: str,
    body: str,
    message: str,
) -> None:
    path = tmp_path / f"{limit_name}.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, limit_name, 1)
    rewrite_xlsx(path, {"xl/worksheets/sheet1.xml": _worksheet(body)})

    with pytest.raises(LimitExceededError, match=message):
        preflight_xlsx(path)


def test_loader_collection_boundary_values_are_accepted(tmp_path: Path) -> None:
    path = tmp_path / "loader-boundaries.xlsx"
    _one_sheet(path)
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1:B2"/><cols><col min="1" max="1" width="10"/></cols>'
                '<sheetData><row r="1" hidden="1"><c r="A1"><v>1</v></c></row></sheetData>'
                '<mergeCells><mergeCell ref="A1:B1"/></mergeCells>'
                '<conditionalFormatting sqref="A1"><cfRule type="expression" priority="1"/>'
                '</conditionalFormatting><dataValidations count="1">'
                '<dataValidation type="whole" sqref="A1"/></dataValidations>'
                '<rowBreaks count="1"><brk id="1"/></rowBreaks>'
                '<scenarios><scenario name="one"/></scenarios>'
            )
        },
    )

    usage = preflight_xlsx(path).usage

    assert usage.serialized_cells == 1
    assert usage.non_empty_cells == 1
    assert usage.merge_ranges == 1
    assert usage.merge_footprint == 2
    assert usage.conditional_formatting_rules == 1
    assert usage.data_validations == 1
    assert usage.row_dimensions == 1
    assert usage.column_dimensions == 1
    assert usage.page_breaks == 1
    assert usage.scenarios == 1


def test_shared_string_item_and_text_budgets_have_boundary_and_overflow(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "strings.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_SHARED_STRINGS", 2)
    monkeypatch.setattr(preflight_module, "MAX_SHARED_STRING_CHARS", 3)
    rewrite_xlsx(
        path,
        {
            "xl/sharedStrings.xml": (
                f'<sst xmlns="{SHEET_NS}"><si><t>a</t></si><si><t>bc</t></si></sst>'
            ).encode()
        },
    )
    assert preflight_xlsx(path).usage.shared_strings == 2

    rewrite_xlsx(
        path,
        {"xl/sharedStrings.xml": (f'<sst xmlns="{SHEET_NS}"><si><t>abcd</t></si></sst>').encode()},
    )
    with pytest.raises(LimitExceededError, match="shared string text"):
        preflight_xlsx(path)


@pytest.mark.parametrize(
    ("element", "limit_name", "message"),
    [
        ("font", "MAX_FONTS", "font"),
        ("fill", "MAX_FILLS", "fill"),
        ("border", "MAX_BORDERS", "border"),
        ("xf", "MAX_CELL_XFS", "cellXfs"),
        ("dxf", "MAX_DXFS", "dxf"),
    ],
)
def test_stylesheet_loader_collections_are_bounded(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    element: str,
    limit_name: str,
    message: str,
) -> None:
    path = tmp_path / f"{element}.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, limit_name, 1)
    container = {
        "font": "fonts",
        "fill": "fills",
        "border": "borders",
        "xf": "cellXfs",
        "dxf": "dxfs",
    }[element]
    rewrite_xlsx(
        path,
        {
            "xl/styles.xml": (
                f'<styleSheet xmlns="{SHEET_NS}"><{container} count="2">'
                f"<{element}/><{element}/></{container}></styleSheet>"
            ).encode()
        },
    )

    with pytest.raises(LimitExceededError, match=message):
        preflight_xlsx(path)


def test_stylesheet_boundary_values_are_counted(tmp_path: Path) -> None:
    path = tmp_path / "styles-boundary.xlsx"
    _one_sheet(path)
    rewrite_xlsx(
        path,
        {
            "xl/styles.xml": (
                f'<styleSheet xmlns="{SHEET_NS}"><fonts><font/></fonts><fills><fill/></fills>'
                "<borders><border/></borders><cellStyleXfs><xf/></cellStyleXfs>"
                "<cellXfs><xf/></cellXfs><cellStyles><cellStyle/></cellStyles>"
                "<dxfs><dxf/></dxfs><tableStyles><tableStyle/></tableStyles></styleSheet>"
            ).encode()
        },
    )

    usage = preflight_xlsx(path).usage

    assert usage.style_records == 8
    assert usage.cell_xfs == 1
    assert usage.dxfs == 1


def test_defined_names_and_custom_properties_are_bounded(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "names.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_DEFINED_NAMES", 1)
    one_name = (
        f'<workbook xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}"><sheets>'
        '<sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets><definedNames>'
        '<definedName name="one">Sheet!$A$1</definedName></definedNames></workbook>'
    ).encode()
    rewrite_xlsx(path, {"xl/workbook.xml": one_name})
    assert preflight_xlsx(path).usage.defined_names == 1

    rewrite_xlsx(
        path,
        {
            "xl/workbook.xml": (
                f'<workbook xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}"><sheets>'
                '<sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets><definedNames>'
                '<definedName name="one">Sheet!$A$1</definedName>'
                '<definedName name="two">Sheet!$A$2</definedName>'
                "</definedNames></workbook>"
            ).encode()
        },
    )
    with pytest.raises(LimitExceededError, match="defined name"):
        preflight_xlsx(path)

    monkeypatch.setattr(preflight_module, "MAX_DEFINED_NAMES", 10)
    monkeypatch.setattr(preflight_module, "MAX_CUSTOM_PROPERTIES", 1)
    rewrite_xlsx(
        path,
        {
            "docProps/custom.xml": (
                b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
                b'custom-properties"><property name="one"/></Properties>'
            )
        },
    )
    assert preflight_xlsx(path).usage.custom_properties == 1

    rewrite_xlsx(
        path,
        {
            "docProps/custom.xml": (
                b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
                b'custom-properties"><property name="one"/><property name="two"/></Properties>'
            )
        },
    )
    with pytest.raises(LimitExceededError, match="custom propert"):
        preflight_xlsx(path)


def test_table_count_and_footprint_are_bounded_before_loader(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "table.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_TABLES", 1)
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1:D2"/><sheetData/><tableParts count="2">'
                '<tablePart r:id="rIdT1"/><tablePart r:id="rIdT2"/></tableParts>'
            ),
            "xl/worksheets/_rels/sheet1.xml.rels": _relationships(
                ("rIdT1", f"{OFFICE_REL_NS}/table", "../tables/table1.xml"),
                ("rIdT2", f"{OFFICE_REL_NS}/table", "../tables/table2.xml"),
            ),
            "xl/tables/table1.xml": f'<table xmlns="{SHEET_NS}" ref="A1:B2"/>'.encode(),
            "xl/tables/table2.xml": f'<table xmlns="{SHEET_NS}" ref="C1:D2"/>'.encode(),
        },
    )
    with pytest.raises(LimitExceededError, match="table count"):
        preflight_xlsx(path)

    monkeypatch.setattr(preflight_module, "MAX_TABLES", 2)
    monkeypatch.setattr(preflight_module, "MAX_TABLE_FOOTPRINT", 8)
    assert preflight_xlsx(path).usage.tables == 2

    monkeypatch.setattr(preflight_module, "MAX_TABLE_FOOTPRINT", 7)
    with pytest.raises(LimitExceededError, match="table footprint"):
        preflight_xlsx(path)


def test_drawing_object_chart_cache_and_anchor_are_preflighted(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "drawing.xlsx"
    _one_sheet(path)
    drawing_rel = f"{OFFICE_REL_NS}/drawing"
    chart_rel = f"{OFFICE_REL_NS}/chart"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    drawing = (
        f'<xdr:wsDr xmlns:xdr="{drawing_ns}" xmlns:c="{chart_ns}" '
        f'xmlns:r="{OFFICE_REL_NS}"><xdr:oneCellAnchor><xdr:from>'
        "<xdr:col>0</xdr:col><xdr:row>0</xdr:row></xdr:from>"
        '<xdr:ext cx="1" cy="1"/><xdr:graphicFrame><c:chart r:id="rIdC1"/>'
        "</xdr:graphicFrame><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>"
    ).encode()
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1"/><sheetData/><drawing r:id="rIdD1"/>'
            ),
            "xl/worksheets/_rels/sheet1.xml.rels": _relationships(
                ("rIdD1", drawing_rel, "../drawings/drawing1.xml"),
            ),
            "xl/drawings/drawing1.xml": drawing,
            "xl/drawings/_rels/drawing1.xml.rels": _relationships(
                ("rIdC1", chart_rel, "../charts/chart1.xml"),
            ),
            "xl/charts/chart1.xml": (
                f'<c:chartSpace xmlns:c="{chart_ns}"><c:numCache><c:pt idx="0"><c:v>1</c:v>'
                '</c:pt><c:pt idx="1"><c:v>2</c:v></c:pt></c:numCache></c:chartSpace>'
            ).encode(),
        },
    )
    monkeypatch.setattr(preflight_module, "MAX_CHART_CACHE_POINTS", 1)
    with pytest.raises(LimitExceededError, match="chart cache"):
        preflight_xlsx(path)

    monkeypatch.setattr(preflight_module, "MAX_CHART_CACHE_POINTS", 2)
    assert preflight_xlsx(path).usage.drawing_objects == 1

    monkeypatch.setattr(preflight_module, "MAX_DRAWING_OBJECTS", 0)
    with pytest.raises(LimitExceededError, match="drawing object"):
        preflight_xlsx(path)
    monkeypatch.setattr(preflight_module, "MAX_DRAWING_OBJECTS", 256)

    invalid = drawing.replace(b"<xdr:col>0</xdr:col>", b"<xdr:col>16384</xdr:col>")
    rewrite_xlsx(path, {"xl/drawings/drawing1.xml": invalid})
    with pytest.raises(CorruptDocumentError, match="drawing anchor"):
        preflight_xlsx(path)


def test_hyperlink_comment_and_relationship_references_are_bounded_and_resolved(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "links.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_HYPERLINKS_AND_COMMENTS", 1)
    one_link = _worksheet(
        '<dimension ref="A1"/><sheetData/><hyperlinks>'
        '<hyperlink ref="A1" r:id="rIdH1"/></hyperlinks>'
    )
    one_link_relationship = (
        f'<Relationships xmlns="{PACKAGE_REL_NS}">'
        f'<Relationship Id="rIdH1" Type="{OFFICE_REL_NS}/hyperlink" '
        'Target="https://example.com/1" TargetMode="External"/></Relationships>'
    ).encode()
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": one_link,
            "xl/worksheets/_rels/sheet1.xml.rels": one_link_relationship,
        },
    )
    assert preflight_xlsx(path).usage.hyperlinks_and_comments == 1

    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1:B1"/><sheetData/><hyperlinks>'
                '<hyperlink ref="A1" r:id="rIdH1"/><hyperlink ref="B1" r:id="rIdH2"/>'
                "</hyperlinks>"
            ),
            "xl/worksheets/_rels/sheet1.xml.rels": (
                f'<Relationships xmlns="{PACKAGE_REL_NS}">'
                f'<Relationship Id="rIdH1" Type="{OFFICE_REL_NS}/hyperlink" '
                'Target="https://example.com/1" TargetMode="External"/>'
                f'<Relationship Id="rIdH2" Type="{OFFICE_REL_NS}/hyperlink" '
                'Target="https://example.com/2" TargetMode="External"/></Relationships>'
            ).encode(),
        },
    )
    with pytest.raises(LimitExceededError, match="hyperlink and comment"):
        preflight_xlsx(path)

    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1"/><sheetData/><drawing r:id="missing"/>'
            ),
        },
    )
    with pytest.raises(CorruptDocumentError, match="relationship"):
        preflight_xlsx(path)


def test_pivot_cache_collections_are_bounded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "pivot.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_PIVOT_CACHES", 1)
    one_cache_workbook = (
        f'<workbook xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}"><sheets>'
        '<sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets><pivotCaches>'
        '<pivotCache cacheId="1" r:id="rIdP1"/></pivotCaches></workbook>'
    ).encode()
    one_cache_relationships = _relationships(
        ("rId1", f"{OFFICE_REL_NS}/worksheet", "worksheets/sheet1.xml"),
        ("rIdP1", f"{OFFICE_REL_NS}/pivotCacheDefinition", "pivotCache/a.xml"),
    )
    rewrite_xlsx(
        path,
        {
            "xl/workbook.xml": one_cache_workbook,
            "xl/_rels/workbook.xml.rels": one_cache_relationships,
            "xl/pivotCache/a.xml": (f'<pivotCacheDefinition xmlns="{SHEET_NS}"/>').encode(),
        },
    )
    assert preflight_xlsx(path).usage.pivot_caches == 1

    rewrite_xlsx(
        path,
        {
            "xl/workbook.xml": (
                f'<workbook xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}"><sheets>'
                '<sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets><pivotCaches>'
                '<pivotCache cacheId="1" r:id="rIdP1"/><pivotCache cacheId="2" r:id="rIdP2"/>'
                "</pivotCaches></workbook>"
            ).encode(),
            "xl/_rels/workbook.xml.rels": _relationships(
                ("rId1", f"{OFFICE_REL_NS}/worksheet", "worksheets/sheet1.xml"),
                ("rIdP1", f"{OFFICE_REL_NS}/pivotCacheDefinition", "pivotCache/a.xml"),
                ("rIdP2", f"{OFFICE_REL_NS}/pivotCacheDefinition", "pivotCache/b.xml"),
            ),
            "xl/pivotCache/a.xml": b"<pivotCacheDefinition/>",
            "xl/pivotCache/b.xml": b"<pivotCacheDefinition/>",
        },
    )
    with pytest.raises(LimitExceededError, match="pivot cache"):
        preflight_xlsx(path)


def test_comment_and_pivot_record_limits_apply_before_loader(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "comments-pivot.xlsx"
    _one_sheet(path)
    monkeypatch.setattr(preflight_module, "MAX_HYPERLINKS_AND_COMMENTS", 1)
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/_rels/sheet1.xml.rels": _relationships(
                ("rIdC1", f"{OFFICE_REL_NS}/comments", "../comments1.xml"),
            ),
            "xl/comments1.xml": (
                f'<comments xmlns="{SHEET_NS}"><authors><author>A</author></authors>'
                '<commentList><comment ref="A1"><text><t>one</t></text></comment>'
                "</commentList></comments>"
            ).encode(),
        },
    )
    assert preflight_xlsx(path).usage.hyperlinks_and_comments == 1

    rewrite_xlsx(
        path,
        {
            "xl/comments1.xml": (
                f'<comments xmlns="{SHEET_NS}"><authors><author>A</author></authors>'
                '<commentList><comment ref="A1"><text><t>one</t></text></comment>'
                '<comment ref="A2"><text><t>two</t></text></comment></commentList></comments>'
            ).encode()
        },
    )
    with pytest.raises(LimitExceededError, match="hyperlink and comment"):
        preflight_xlsx(path)

    monkeypatch.setattr(preflight_module, "MAX_HYPERLINKS_AND_COMMENTS", 20_000)
    monkeypatch.setattr(preflight_module, "MAX_PIVOT_CACHE_RECORDS", 1)
    rewrite_xlsx(
        path,
        {
            "xl/comments1.xml": None,
            "xl/worksheets/_rels/sheet1.xml.rels": None,
            "xl/pivotCache/cache.xml": f'<pivotCacheDefinition xmlns="{SHEET_NS}"/>'.encode(),
            "xl/pivotCache/_rels/cache.xml.rels": _relationships(
                (
                    "rIdR1",
                    f"{OFFICE_REL_NS}/pivotCacheRecords",
                    "records.xml",
                ),
            ),
            "xl/pivotCache/records.xml": (
                f'<pivotCacheRecords xmlns="{SHEET_NS}"><r/><r/></pivotCacheRecords>'
            ).encode(),
            "xl/workbook.xml": (
                f'<workbook xmlns="{SHEET_NS}" xmlns:r="{OFFICE_REL_NS}"><sheets>'
                '<sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets><pivotCaches>'
                '<pivotCache cacheId="1" r:id="rIdP1"/></pivotCaches></workbook>'
            ).encode(),
            "xl/_rels/workbook.xml.rels": _relationships(
                ("rId1", f"{OFFICE_REL_NS}/worksheet", "worksheets/sheet1.xml"),
                (
                    "rIdP1",
                    f"{OFFICE_REL_NS}/pivotCacheDefinition",
                    "pivotCache/cache.xml",
                ),
            ),
        },
    )
    with pytest.raises(LimitExceededError, match="pivot cache record"):
        preflight_xlsx(path)


def test_unknown_relationship_objects_are_deterministically_locatable(tmp_path: Path) -> None:
    path = tmp_path / "unknown.xlsx"
    _one_sheet(path)
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/_rels/sheet1.xml.rels": _relationships(
                ("rIdOle", f"{OFFICE_REL_NS}/oleObject", "../embeddings/ole1.bin"),
            ),
            "xl/embeddings/ole1.bin": b"opaque",
        },
    )

    unsupported = preflight_xlsx(path).sheets[0].unsupported_objects

    assert [
        (item.source_index, item.kind, item.relationship_id, item.target) for item in unsupported
    ] == [(0, "oleObject", "rIdOle", "xl/embeddings/ole1.bin")]


@pytest.mark.parametrize(
    "xml",
    [
        b'<!DOCTYPE worksheet [<!ENTITY x "boom">]><worksheet xmlns="http://schemas.'
        b'openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>',
        b'<worksheet xmlns="urn:not-spreadsheet"><sheetData/></worksheet>',
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
    ],
)
def test_worksheet_dtd_entity_namespace_and_malformed_xml_are_typed_corruption(
    tmp_path: Path,
    xml: bytes,
) -> None:
    path = tmp_path / "corrupt.xlsx"
    _one_sheet(path)
    rewrite_xlsx(path, {"xl/worksheets/sheet1.xml": xml})

    with pytest.raises(CorruptDocumentError):
        preflight_xlsx(path)


def test_zip_slip_is_rejected_by_preflight(tmp_path: Path) -> None:
    path = tmp_path / "zip-slip.xlsx"
    _one_sheet(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("../escape.xml", b"<x/>")

    with pytest.raises(CorruptDocumentError, match="unsafe member"):
        preflight_xlsx(path)


def test_projected_wire_budget_is_a_stricter_success_boundary_than_grid_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "wire.xlsx"
    _one_sheet(path, cells=("A1", "B1"))
    assert preflight_module.MAX_MATERIALIZED_GRID_CELLS == 200_000
    monkeypatch.setattr(preflight_module, "MAX_PROJECTED_WIRE_BYTES", 2_000)

    with pytest.raises(LimitExceededError, match="inline result budget"):
        preflight_xlsx(path)


def test_projected_workbook_budget_bounds_full_mode_loader_work(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "projected-workbook.xlsx"
    _one_sheet(path, cells=("A1",))
    monkeypatch.setattr(preflight_module, "MAX_PROJECTED_WORKBOOK_BYTES", 1)

    with pytest.raises(LimitExceededError, match="projected workbook"):
        preflight_xlsx(path)


def test_materialized_grid_and_native_text_have_independent_outer_limits(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "grid-text.xlsx"
    _one_sheet(path, cells=("A1", "B1"))
    monkeypatch.setattr(preflight_module, "MAX_MATERIALIZED_GRID_CELLS", 1)
    with pytest.raises(LimitExceededError, match="materialized grid"):
        preflight_xlsx(path)

    monkeypatch.setattr(preflight_module, "MAX_MATERIALIZED_GRID_CELLS", 200_000)
    monkeypatch.setattr(preflight_module, "MAX_NATIVE_TEXT_CHARS", 3)
    rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": _worksheet(
                '<dimension ref="A1"/><sheetData><row><c r="A1"><v>1234</v></c></row></sheetData>'
            )
        },
    )
    with pytest.raises(LimitExceededError, match="native text"):
        preflight_xlsx(path)
