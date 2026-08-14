from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl.styles.numbers import is_date_format, is_timedelta_format
from openpyxl.utils.datetime import from_excel

_CURRENCY_SYMBOLS = frozenset({"$", "€", "£", "¥"})
_BRACKET_TOKEN = re.compile(r"\[([^]]+)]")
_SCIENTIFIC_FORMAT = re.compile(r"[0#?](?:\.[0#?]+)?[Ee][+-]?[0#?]+")
_FRACTION_FORMAT = re.compile(r"(?:^|[^A-Za-z])[0#?]+(?:\s+[0#?]+)?/[0#?]+")
_LOCALE_CURRENCY = re.compile(r"^\$([^\]-]*)-[0-9A-Fa-f]+$")
_QUOTED_LITERAL = re.compile(r'"([^"]*)"')


@dataclass(frozen=True, slots=True)
class FormattedSavedValue:
    text: str
    warning: str | None = None


def _stable_decimal(value: Decimal) -> str:
    if not value.is_finite():
        return str(value)
    rendered = format(value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered or "0"


def stable_raw_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, timedelta):
        return _stable_decimal(Decimal(str(value.total_seconds())))
    if isinstance(value, Decimal):
        return _stable_decimal(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return _stable_decimal(Decimal(str(value)))
    return str(value)


def _split_sections(number_format: str) -> tuple[str, ...]:
    sections: list[str] = []
    current: list[str] = []
    quoted = False
    bracket_depth = 0
    escaped = False
    for character in number_format:
        if escaped:
            current.append(character)
            escaped = False
        elif character == "\\":
            current.append(character)
            escaped = True
        elif character == '"':
            current.append(character)
            quoted = not quoted
        elif not quoted and character == "[":
            bracket_depth += 1
            current.append(character)
        elif not quoted and character == "]":
            bracket_depth = max(0, bracket_depth - 1)
            current.append(character)
        elif not quoted and bracket_depth == 0 and character == ";":
            sections.append("".join(current))
            current = []
        else:
            current.append(character)
    sections.append("".join(current))
    return tuple(sections)


def _has_unsupported_token(number_format: str) -> bool:
    if _SCIENTIFIC_FORMAT.search(number_format) or _FRACTION_FORMAT.search(number_format):
        return True
    for literal in _QUOTED_LITERAL.findall(number_format):
        if any(character not in {*_CURRENCY_SYMBOLS, " ", "-", "(", ")"} for character in literal):
            return True
    for match in _BRACKET_TOKEN.finditer(number_format):
        token = match.group(1)
        if token.casefold() in {"h", "hh", "m", "mm", "s", "ss"}:
            continue
        currency = _LOCALE_CURRENCY.fullmatch(token)
        if currency is not None and currency.group(1) in _CURRENCY_SYMBOLS:
            continue
        return True
    remaining = _QUOTED_LITERAL.sub("", number_format)
    remaining = _BRACKET_TOKEN.sub("", remaining)
    remaining = remaining.casefold().replace("am/pm", "")
    index = 0
    while index < len(remaining):
        character = remaining[index]
        if character in {"_", "*"}:
            index += 2
            continue
        if character == "\\" and index + 1 < len(remaining):
            if remaining[index + 1].isalnum():
                return True
            index += 2
            continue
        if character.isalpha() and character not in "ymdhs":
            return True
        index += 1
    return False


def _clean_section(section: str) -> str:
    cleaned: list[str] = []
    index = 0
    while index < len(section):
        character = section[index]
        if character in {"_", "*"}:
            index += 2
            continue
        if character == "\\" and index + 1 < len(section):
            cleaned.append(section[index + 1])
            index += 2
            continue
        if character == '"':
            end = section.find('"', index + 1)
            if end == -1:
                return section
            cleaned.append(section[index + 1 : end])
            index = end + 1
            continue
        if character == "[":
            end = section.find("]", index + 1)
            if end == -1:
                return section
            token = section[index + 1 : end]
            currency = _LOCALE_CURRENCY.fullmatch(token)
            cleaned.append(currency.group(1) if currency is not None else f"[{token}]")
            index = end + 1
            continue
        cleaned.append(character)
        index += 1
    return "".join(cleaned).strip()


def _as_decimal(value: object) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float) and math.isfinite(value):
        return Decimal(str(value))
    return None


def _selected_numeric_section(sections: tuple[str, ...], value: Decimal) -> tuple[str, bool]:
    if value < 0:
        return (sections[1] if len(sections) > 1 else sections[0], True)
    if value == 0 and len(sections) > 2:
        return sections[2], False
    return sections[0], False


def _decimal_places(pattern: str) -> tuple[int, int]:
    if "." not in pattern:
        return 0, 0
    suffix = pattern.split(".", 1)[1]
    placeholders: list[str] = []
    for character in suffix:
        if character in "0#?":
            placeholders.append(character)
        elif placeholders:
            break
    return placeholders.count("0"), len(placeholders)


def _format_number(value: Decimal, number_format: str) -> str | None:
    sections = _split_sections(number_format)
    section, negative = _selected_numeric_section(sections, value)
    pattern = _clean_section(section)
    if not any(character in "0#?" for character in pattern):
        if value == 0 and "-" in pattern:
            symbol = next((item for item in pattern if item in _CURRENCY_SYMBOLS), "")
            return f"{symbol}-"
        return None

    percent_count = pattern.count("%")
    magnitude = abs(value) * (Decimal(100) ** percent_count)
    integer_pattern = pattern.split(".", 1)[0]
    scale_commas = len(integer_pattern) - len(integer_pattern.rstrip(","))
    if scale_commas:
        magnitude /= Decimal(1000) ** scale_commas
        integer_pattern = integer_pattern.rstrip(",")
    minimum_decimals, maximum_decimals = _decimal_places(pattern)
    if maximum_decimals:
        quantum = Decimal(1).scaleb(-maximum_decimals)
        magnitude = magnitude.quantize(quantum, rounding=ROUND_HALF_UP)
        rendered = f"{magnitude:.{maximum_decimals}f}"
        integer, fraction = rendered.split(".", 1)
        if maximum_decimals > minimum_decimals:
            fraction = fraction.rstrip("0")
            fraction += "0" * max(0, minimum_decimals - len(fraction))
        rendered = integer + (f".{fraction}" if fraction else "")
    else:
        rendered = str(magnitude.quantize(Decimal(1), rounding=ROUND_HALF_UP))

    integer, separator, fraction = rendered.partition(".")
    if "," in integer_pattern:
        integer = f"{int(integer):,}"
    rendered = integer + (separator + fraction if separator else "")

    currency = next((symbol for symbol in _CURRENCY_SYMBOLS if symbol in pattern), "")
    first_placeholder = min(
        (pattern.find(character) for character in "0#?" if character in pattern),
        default=0,
    )
    if currency:
        rendered = (
            f"{currency}{rendered}"
            if pattern.find(currency) <= first_placeholder
            else f"{rendered}{currency}"
        )
    if percent_count:
        rendered += "%" * percent_count
    if negative:
        rendered = f"({rendered})" if "(" in pattern and ")" in pattern else f"-{rendered}"
    return rendered


def _time_fraction_digits(number_format: str) -> int:
    match = re.search(r"s{1,2}\.([0#]+)", number_format, flags=re.IGNORECASE)
    return len(match.group(1)) if match is not None else 0


def _render_clock(value: time, number_format: str) -> str:
    lowered = number_format.casefold()
    include_seconds = "s" in lowered
    twelve_hour = "am/pm" in lowered
    raw_hour = value.hour % 12 or 12 if twelve_hour else value.hour
    hour_token = re.search(r"(?<!\[)(h{1,2})", lowered)
    minute_token = re.search(r"(?<=:)(m{1,2})", lowered)
    if hour_token is None:
        minute_only = re.fullmatch(r"m{1,2}:s{1,2}(?:\.[0#]+)?", lowered)
        if minute_only is not None:
            return f"{value.minute:02d}:{value.second:02d}"
        if re.fullmatch(r"s{1,2}(?:\.[0#]+)?", lowered):
            return f"{value.second:02d}"
    hour = (
        f"{raw_hour:02d}"
        if hour_token is not None and len(hour_token.group(1)) == 2
        else str(raw_hour)
    )
    minute = (
        f"{value.minute:02d}"
        if minute_token is None or len(minute_token.group(1)) == 2
        else str(value.minute)
    )
    base = f"{hour}:{minute}"
    if include_seconds:
        base += f":{value.second:02d}"
    digits = _time_fraction_digits(number_format)
    if digits:
        fraction = f"{value.microsecond:06d}"[:digits].ljust(digits, "0")
        base += f".{fraction}"
    if twelve_hour:
        base += " AM" if value.hour < 12 else " PM"
    return base


def _render_elapsed(value: timedelta, number_format: str) -> str:
    total_seconds = Decimal(str(value.total_seconds()))
    digits = _time_fraction_digits(number_format)
    quantum = Decimal(1).scaleb(-digits) if digits else Decimal(1)
    total_seconds = total_seconds.quantize(quantum, rounding=ROUND_HALF_UP)
    negative = total_seconds < 0
    total_seconds = abs(total_seconds)
    whole_seconds = int(total_seconds)
    fraction = total_seconds - whole_seconds
    lowered = number_format.casefold()
    if "[h]" in lowered or "[hh]" in lowered:
        total_hours = whole_seconds // 3600
        rendered_hours = f"{total_hours:02d}" if "[hh]" in lowered else str(total_hours)
        rendered = f"{rendered_hours}:{whole_seconds % 3600 // 60:02d}"
        if "s" in lowered:
            rendered += f":{whole_seconds % 60:02d}"
    elif "[m]" in lowered or "[mm]" in lowered:
        total_minutes = whole_seconds // 60
        rendered = f"{total_minutes:02d}" if "[mm]" in lowered else str(total_minutes)
        if "s" in lowered:
            rendered += f":{whole_seconds % 60:02d}"
    else:
        rendered = str(whole_seconds)
    if digits:
        rendered += f".{str(fraction)[2:]:0<{digits}}"[: digits + 1]
    return f"-{rendered}" if negative else rendered


def _format_date_value(value: object, number_format: str, epoch: datetime) -> str | None:
    converted = value
    decimal = _as_decimal(value)
    try:
        if decimal is not None:
            converted = from_excel(
                float(decimal),
                epoch,
                timedelta=is_timedelta_format(number_format),
            )
    except (OverflowError, ValueError):
        return None
    if isinstance(converted, timedelta):
        return _render_elapsed(converted, number_format)
    if isinstance(converted, datetime):
        lowered = number_format.casefold()
        has_date = "y" in lowered or "d" in lowered
        has_time = "h" in lowered or "s" in lowered or not has_date
        if has_date and has_time:
            return (
                f"{converted.date().isoformat()} {_render_clock(converted.time(), number_format)}"
            )
        if has_date:
            return converted.date().isoformat()
        return _render_clock(converted.time(), number_format)
    if isinstance(converted, date):
        return converted.isoformat()
    if isinstance(converted, time):
        return _render_clock(converted, number_format)
    return None


def format_saved_value(
    value: object,
    number_format: str,
    *,
    epoch: datetime,
    conditional_number_format: bool = False,
) -> FormattedSavedValue:
    raw = stable_raw_value(value)
    if value is None or isinstance(value, bool | str):
        return FormattedSavedValue(raw)
    if conditional_number_format:
        return FormattedSavedValue(raw, "unsupported number format")
    if number_format.casefold() == "general":
        return FormattedSavedValue(raw)
    if _has_unsupported_token(number_format):
        return FormattedSavedValue(raw, "unsupported number format")
    if is_date_format(number_format):
        rendered_date = _format_date_value(value, number_format, epoch)
        if rendered_date is not None:
            return FormattedSavedValue(rendered_date)
        return FormattedSavedValue(raw, "unsupported number format")
    decimal = _as_decimal(value)
    if decimal is None:
        return FormattedSavedValue(raw)
    try:
        rendered_number = _format_number(decimal, number_format)
    except (InvalidOperation, ValueError):
        rendered_number = None
    if rendered_number is None:
        return FormattedSavedValue(raw, "unsupported number format")
    return FormattedSavedValue(rendered_number)
